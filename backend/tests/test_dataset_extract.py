"""Unit tests for dataset file text extraction."""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from openpyxl import Workbook

from app.dataset.rag.extract import extract_text_from_bytes


def _docx_bytes(*paragraphs: str) -> bytes:
    """Build an in-memory DOCX payload for extraction tests."""

    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    """Build an in-memory XLSX payload for extraction tests."""

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_docx_from_bytes_without_temp_file() -> None:
    """DOCX extraction reads directly from bytes (Windows-safe)."""

    payload = _docx_bytes("paragraph one", "paragraph two")
    text = extract_text_from_bytes(payload, file_name="sample.docx")
    assert "paragraph one" in text
    assert "paragraph two" in text


def test_extract_xlsx_from_bytes_without_temp_file() -> None:
    """XLSX extraction reads directly from bytes (Windows-safe)."""

    payload = _xlsx_bytes([["A", "B"], ["1", "2"]])
    text = extract_text_from_bytes(payload, file_name="sample.xlsx")
    assert "A" in text
    assert "2" in text
